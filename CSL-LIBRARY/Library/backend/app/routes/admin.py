import io
import openpyxl
from flask import Blueprint, render_template, request, redirect, url_for, send_file, jsonify,g
from app.utils import admin_required
from app.models import *
from app.extension import db
from datetime import datetime, timedelta
from werkzeug.utils import secure_filename
import os

import re as _re
from collections import OrderedDict as _OD


def _base_code(code: str) -> str:
    return _re.sub(r"-c\.\d+$", "", code)

admin_bp  = Blueprint("admin", __name__, url_prefix="/admin")
BASE_DIR = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..", "..")
)
COVER_DIR = os.path.join(
    BASE_DIR,
    "frontend",
    "storage",
    "Book-cover"
)


# ─────────────────────────────────────────────
#  SHARED HELPERS
# ─────────────────────────────────────────────

def get_next_running_no(main_name: str, subject_name: str) -> int:
    """
    Finds the highest existing running number for a given main+subject
    combination by scanning existing book codes, then returns next integer.

    Book code format: MAIN-SUBJECT-001-c.N
    e.g. TXT-NAV-001-c.1, TXT-NAV-001-c.2, TXT-NAV-002-c.1 ...
    """
    prefix = f"{main_name}-{subject_name}-"
    books  = Book.query.filter(Book.code.ilike(f"{prefix}%")).all()

    if not books:
        return 1

    max_num = 0
    for book in books:
        suffix = book.code[len(prefix):]       # e.g. "003-c.1"
        try:
            num = int(suffix.split("-")[0])    # grab the running number part
            if num > max_num:
                max_num = num
        except ValueError:
            continue

    return max_num + 1


def build_code(main_name: str, subject_name: str, running_no: int, copy_no: str) -> str:
    """Builds 4-part code: MAIN-SUBJECT-001-c.1"""
    return f"{main_name}-{subject_name}-{str(running_no).zfill(3)}-{copy_no}"


def is_code_unique(code: str) -> bool:
    return Book.query.filter_by(code=code).first() is None


def get_employee_indicator(employee):
    now    = datetime.utcnow()
    active = [r for r in employee.records if r.status == "APPROVED"]
    if not active:
        return "green"
    if any(r.expected_at and r.expected_at < now for r in active):
        return "red"
    return "orange"


# ═════════════════════════════════════════════
#  DASHBOARD
# ═════════════════════════════════════════════

@admin_bp.route("/")
@admin_required
def dashboard():
    now = datetime.utcnow()

    books_count     = Book.query.count()
    employees_count = Employee.query.filter_by(status="ACTIVE").count()
    pending_count   = Record.query.filter_by(status="PENDING").count()
    overdue         = Record.query.filter(
                          Record.status == "APPROVED",
                          Record.expected_at < now
                      ).count()

    return render_template(
        "admin/dashboard.html",
        books_count     = books_count,
        employees_count = employees_count,
        pending_count   = pending_count,
        overdue         = overdue,
    )


# ══════════════════════════════════════════════
#  BOOKS SECTION
# ══════════════════════════════════════════════

@admin_bp.route("/books")
@admin_required
def books():
    search     = request.args.get("search",     "").strip()
    search_by  = request.args.get("search_by",  "title")
    subject_id = request.args.get("subject_id", "").strip()
    main_id    = request.args.get("main_id",    "").strip()
    available  = request.args.get("available",  "").strip()
    page       = request.args.get("page", 1, type=int)

    q = Book.query

    if search:
        if search_by == "title":
            q = q.filter(Book.title.ilike(f"%{search}%"))
        elif search_by == "code":
            q = q.filter(Book.code.ilike(f"%{search}%"))
        elif search_by == "author":
            q = q.filter(Book.author.ilike(f"%{search}%"))
        elif search_by == "keyword":
            q = q.filter(Book.keywords.ilike(f"%{search}%"))

    if subject_id:
        q = q.filter(Book.subject_category == int(subject_id))

    if main_id:
        main = Main.query.get(int(main_id))
        if main:
            q = q.filter(Book.code.ilike(f"{main.name}-%"))

    if available == "1":
        q = q.filter(Book.available_copies > 0)

    # Paginate over individual Book rows
    books_page = q.order_by(Book.code).paginate(page=page, per_page=24, error_out=False)

    # ── Group the page's items by base running code ───────────────────────────
    groups: _OD = _OD()
    for book in books_page.items:
        base = _base_code(book.code)
        if base not in groups:
            groups[base] = {
                "base_code":    base,
                "rep":          book,
                "total_copies": 0,
                "avail_copies": 0,
                "codes":        [],
            }
        g = groups[base]
        g["total_copies"] += 1
        g["avail_copies"] += (book.available_copies or 0)
        g["codes"].append(book.code)

    grouped_books  = list(groups.values())
    grouped_count  = len(grouped_books)

    mains    = Main.query.order_by(Main.name).all()
    subjects = Subject.query.order_by(Subject.name).all()

    return render_template(
        "admin/books/index.html",
        books         = books_page,
        grouped_books = grouped_books,
        grouped_count = grouped_count,
        mains         = mains,
        subjects      = subjects,
        search        = search,
        search_by     = search_by,
        main_id       = main_id,
        subject_id    = subject_id,
        available     = available,
    )

# ── AJAX: live code preview ───────────────────

@admin_bp.route("/books/check-code")
@admin_required
def check_code():
    main    = request.args.get("main",    "").strip().upper()
    subject = request.args.get("subject", "").strip().upper()
    running = request.args.get("running", "").strip()
    copies  = request.args.get("copies",  "1").strip()

    if not all([main, subject]):
        return jsonify({"ready": False})

    try:
        run_no     = int(running) if running else get_next_running_no(main, subject)
        num_copies = max(1, int(copies)) if copies.isdigit() else 1
        codes      = [build_code(main, subject, run_no, f"c.{i}")
                      for i in range(1, num_copies + 1)]
        unique     = all(is_code_unique(c) for c in codes)
        return jsonify({
            "ready":  True,
            "codes":  codes,
            "code":   codes[0],
            "unique": unique,
            "run_no": run_no,
            "auto":   not bool(running),
        })
    except ValueError:
        return jsonify({"ready": False})


# ── ADD BOOK ──────────────────────────────────

@admin_bp.route("/books/add", methods=["GET", "POST"])
@admin_required
def add_book():
    mains    = Main.query.order_by(Main.name).all()
    subjects = Subject.query.order_by(Subject.name).all()
    shelves  = Shelf.query.order_by(Shelf.name).all()

    if request.method == "POST":
        mode = request.form.get("mode")

        # ══════════════════════════
        #  MANUAL
        # ══════════════════════════
        if mode == "manual":
            main_id    = request.form.get("main_id",    type=int)
            subject_id = request.form.get("subject_id", type=int)
            title      = request.form.get("title",      "").strip()
            running_raw = request.form.get("running_no", "").strip()

            missing = []
            if not main_id:    missing.append("Main Category")
            if not subject_id: missing.append("Subject Category")
            if not title:      missing.append("Title")
            if missing:
                return render_template("admin/books/add.html",
                                       mains=mains, subjects=subjects, shelves=shelves,
                                       error=f"Required: {', '.join(missing)}")

            main    = Main.query.get(main_id)
            subject = Subject.query.get(subject_id)

            if running_raw:
                try:
                    running_no = int(running_raw)
                    if running_no < 1:
                        raise ValueError
                except ValueError:
                    return render_template("admin/books/add.html",
                                           mains=mains, subjects=subjects, shelves=shelves,
                                           error="Running number must be a positive integer.")
            else:
                running_no = get_next_running_no(main.name, subject.name)

            num_copies = request.form.get("copies", 1, type=int)
            if num_copies < 1:
                num_copies = 1

            copy_codes = [build_code(main.name, subject.name, running_no, f"c.{i}")
                          for i in range(1, num_copies + 1)]

            for code in copy_codes:
                if not is_code_unique(code):
                    return render_template("admin/books/add.html",
                                           mains=mains, subjects=subjects, shelves=shelves,
                                           error=f"Code '{code}' already exists. "
                                                 f"Choose a different running number.")

            DEFAULT_COVER = "DEFAULTCOVER.png"
            cover_filename = DEFAULT_COVER
            cover_file     = request.files.get("cover_image")
            if cover_file and cover_file.filename:
                cover_filename = secure_filename(cover_file.filename)
                os.makedirs(COVER_DIR, exist_ok=True)
                cover_file.save(os.path.join(COVER_DIR, cover_filename))

            try:
                for copy_code in copy_codes:
                    db.session.add(Book(
                        code             = copy_code,
                        title            = title,
                        author           = request.form.get("author",           "").strip() or None,
                        publisher        = request.form.get("publisher",        "").strip() or None,
                        edition          = request.form.get("edition",          "").strip() or None,
                        publication_year = request.form.get("publication_year", type=int),
                        subject_category = subject_id,
                        shelf_number     = request.form.get("shelf_number",     "").strip() or None,
                        keywords         = request.form.get("keywords",         "").strip() or None,
                        price            = request.form.get("price",            type=int),
                        copies           = num_copies,
                        available_copies = 1,
                        remarks          = request.form.get("remarks",          "").strip() or None,
                        cover_image      = cover_filename,
                    ))
                db.session.commit()
                return redirect(url_for("admin.books"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/add.html",
                                       mains=mains, subjects=subjects, shelves=shelves,
                                       error=f"Failed: {exc}")

        # ══════════════════════════
        #  EXCEL
        # ══════════════════════════
        elif mode == "excel":
            excel_file = request.files.get("excel_file")
            if not excel_file or not excel_file.filename.endswith((".xlsx", ".xls")):
                return render_template("admin/books/add.html",
                                       mains=mains, subjects=subjects, shelves=shelves,
                                       error="Upload a valid .xlsx file.")
            try:
                wb    = openpyxl.load_workbook(io.BytesIO(excel_file.read()), data_only=True)
                sheet = wb.active
            except Exception:
                return render_template("admin/books/add.html",
                                       mains=mains, subjects=subjects, shelves=shelves,
                                       error="Could not read the Excel file.")

            parsed_rows      = []
            all_errors       = []
            skipped          = []
            session_counters = {}

            for row_num, row in enumerate(
                sheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                cols = (list(row) + [None] * 13)[:13]
                if not any(cols):
                    continue

                (main_name, subject_name, running_raw, title, author,
                 publisher, edition, pub_year, shelf_number,
                 keywords, price, copies, remarks) = cols

                main_name    = str(main_name).strip().upper()    if main_name    else None
                subject_name = str(subject_name).strip().upper() if subject_name else None
                running_raw  = str(running_raw).strip()          if running_raw  else None
                title        = str(title).strip()                if title        else None
                author       = str(author).strip()               if author       else None
                publisher    = str(publisher).strip()            if publisher    else None
                edition      = str(edition).strip()              if edition      else None
                pub_year_raw = str(pub_year).strip()             if pub_year     else None
                shelf_number = str(shelf_number).strip()         if shelf_number else None
                copies_raw   = str(copies).strip()               if copies       else None
                keywords     = str(keywords).strip()             if keywords     else None
                remarks      = str(remarks).strip()              if remarks      else None
                price_raw    = str(price).strip()                if price        else None

                row_missing = []
                if not main_name:    row_missing.append("Main")
                if not subject_name: row_missing.append("Subject")
                if not title:        row_missing.append("Title")
                if not copies_raw:   row_missing.append("Copies")

                if row_missing:
                    all_errors.append(
                        f"Row {row_num}: missing — {', '.join(row_missing)}."
                    )
                    continue

                main_obj = Main.query.filter_by(name=main_name).first()
                if not main_obj:
                    all_errors.append(
                        f"Row {row_num}: Main category '{main_name}' not in database. "
                        f"Add it via Code Maker first."
                    )
                    continue

                subj_obj = Subject.query.filter_by(name=subject_name).first()
                if not subj_obj:
                    all_errors.append(
                        f"Row {row_num}: Subject '{subject_name}' not in database. "
                        f"Add it via Code Maker first."
                    )
                    continue

                if shelf_number:
                    shelf_obj = Shelf.query.filter_by(name=shelf_number).first()
                    if not shelf_obj:
                        all_errors.append(
                            f"Row {row_num}: Shelf '{shelf_number}' not in database. "
                            f"Add it via Code Maker first."
                        )
                        continue

                pub_year_int = None
                if pub_year_raw:
                    try:
                        pub_year_int = int(float(pub_year_raw))
                    except (ValueError, TypeError):
                        all_errors.append(f"Row {row_num}: Publication Year must be a number if provided.")
                        continue

                try:
                    num_copies = int(float(copies_raw))
                    if num_copies < 1:
                        raise ValueError
                except (ValueError, TypeError):
                    all_errors.append(f"Row {row_num}: Copies must be a positive integer.")
                    continue

                price_int = None
                if price_raw:
                    try:
                        price_int = int(float(price_raw))
                    except (ValueError, TypeError):
                        all_errors.append(f"Row {row_num}: Price must be a number if provided.")
                        continue

                key = f"{main_name}-{subject_name}"

                if running_raw:
                    try:
                        running_no = int(float(running_raw))
                        if running_no < 1:
                            raise ValueError
                    except (ValueError, TypeError):
                        all_errors.append(
                            f"Row {row_num}: Running No must be a positive integer if provided."
                        )
                        continue
                else:
                    if key in session_counters:
                        running_no = session_counters[key] + 1
                    else:
                        running_no = get_next_running_no(main_name, subject_name)

                session_counters[key] = max(session_counters.get(key, 0), running_no)

                row_codes = [
                    build_code(main_name, subject_name, running_no, f"c.{i}")
                    for i in range(1, num_copies + 1)
                ]

                conflict = None
                for c in row_codes:
                    already_in_buffer = any(r["code"] == c for r in parsed_rows)
                    if not is_code_unique(c) or already_in_buffer:
                        conflict = c
                        break

                if conflict:
                    skipped.append(
                        f"Row {row_num} — {conflict} "
                        f"({'already exists in DB' if not is_code_unique(conflict) else 'duplicate in this file'})"
                    )
                    continue

                for copy_code in row_codes:
                    parsed_rows.append(dict(
                        code             = copy_code,
                        title            = title,
                        author           = author,
                        publisher        = publisher,
                        edition          = edition,
                        publication_year = pub_year_int,
                        subject_category = subj_obj.id,
                        shelf_number     = shelf_number,
                        keywords         = keywords,
                        price            = price_int,
                        copies           = num_copies,
                        available_copies = 1,
                        remarks          = remarks,
                    ))

            if all_errors:
                return render_template("admin/books/add.html",
                                       mains=mains, subjects=subjects, shelves=shelves,
                                       error="File rejected — fix all errors and re-upload. "
                                             "Nothing was imported.",
                                       excel_errors=all_errors,
                                       skipped=skipped)

            try:
                added = []
                for rd in parsed_rows:
                    db.session.add(Book(**rd))
                    added.append(rd["code"])
                db.session.commit()
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/add.html",
                                       mains=mains, subjects=subjects, shelves=shelves,
                                       error=f"Commit failed: {exc}")

            return render_template("admin/books/add.html",
                                   mains=mains, subjects=subjects, shelves=shelves,
                                   success=f"{len(added)} book entr{'y' if len(added)==1 else 'ies'} imported.",
                                   skipped=skipped)

    return render_template("admin/books/add.html", mains=mains, subjects=subjects, shelves=shelves)


# ── EXCEL TEMPLATE DOWNLOAD ───────────────────

@admin_bp.route("/books/excel-template")
@admin_required
def excel_template():
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Books"

    headers = [
        "Main *", "Subject *", "Running No (optional)",
        "Title *", "Author *", "Publisher *", "Edition *",
        "Publication Year *", "Shelf Number *",
        "Keywords", "Price", "Copies *", "Remarks"
    ]
    example = [
        "TXT", "NAV", "",
        "Introduction to Naval Architecture",
        "John Smith", "Elsevier", "3rd", 2020,
        "A1-S3", "naval, ship, design", 450, 2, ""
    ]
    example2 = [
        "TXT", "NAV", 5,
        "Advanced Ship Design",
        "Jane Doe", "Springer", "1st", 2022,
        "A1-S4", "ship, hull, stability", 600, 1, "Reference only"
    ]

    ws.append(headers)
    ws.append(example)
    ws.append(example2)

    for cell in ws[1]:
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")

    optional_cols = [3, 10, 11, 13]
    light_fill = PatternFill("solid", fgColor="D9E1F2")
    for col_idx in optional_cols:
        ws.cell(row=1, column=col_idx).fill = light_fill
        ws.cell(row=1, column=col_idx).font = Font(bold=True, color="333333")

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = (
            max(len(str(c.value or "")) for c in col) + 4)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True, download_name="books_template.xlsx")


# ── EDIT BOOK ─────────────────────────────────
@admin_bp.route("/books/edit", methods=["GET", "POST"])
@admin_required
def edit_book():
    book         = None
    search_error = None
    search_code  = request.args.get("code", "").strip()
    shelves      = Shelf.query.order_by(Shelf.name).all()

    # ── Shared helpers ──────────────────────────────────────────────────
    def get_siblings(code: str):
        base   = _re.sub(r"-c\.\d+$", "", code)
        prefix = f"{base}-c."
        return Book.query.filter(Book.code.ilike(f"{prefix}%")).all()

    def get_available_count(code: str) -> int:
        return sum(s.available_copies for s in get_siblings(code))

    def render_edit(**kwargs):
        b = kwargs.get("book", book)
        kwargs.setdefault("shelves", shelves)
        kwargs.setdefault("available_count", get_available_count(b.code) if b else 0)
        return render_template("admin/books/edit.html", **kwargs)

    # ── Initial GET search ───────────────────────────────────────────────
    if search_code:
        book = Book.query.filter(Book.code.ilike(search_code)).first()
        if not book:
            search_error = f"No book found with code '{search_code}'."

    if request.method == "POST":
        action    = request.form.get("action")
        book_code = request.form.get("book_code", "").strip()
        book      = Book.query.filter_by(code=book_code).first()
        if not book:
            return render_edit(book=None, search_error=f"Book '{book_code}' not found.")

        # ══════════════════════════════════════════════════════════════════
        #  DELETE
        # ══════════════════════════════════════════════════════════════════
        if action == "delete":
            try:
                from app.models import Record as _Record

                # Block deletion if any active/pending borrow records exist
                active_records = _Record.query.filter(
                    _Record.book_id == book.code,
                    _Record.status.in_(["PENDING", "APPROVED", "RETURN_REQUESTED"])
                ).all()

                if active_records:
                    statuses = ", ".join(set(r.status for r in active_records))
                    return render_edit(book=book,
                                        error=f"Cannot delete '{book_code}' — it has active borrow "
                                              f"records ({statuses}). Resolve them first.")

                # Safe to delete — only RETURNED/REJECTED records remain;
                # delete those history records before removing the book
                _Record.query.filter_by(book_id=book.code).delete(synchronize_session=False)
                db.session.delete(book)
                db.session.commit()
                return render_edit(book=None,
                                    success=f"Book '{book_code}' deleted.")
            except Exception as exc:
                db.session.rollback()
                return render_edit(book=book,
                                    error=f"Delete failed: {exc}")

        # ══════════════════════════════════════════════════════════════════
        #  SET AVAILABILITY (across all siblings of the running code)
        # ══════════════════════════════════════════════════════════════════
        elif action == "set_availability":
            avail_val = request.form.get("available_copies", "").strip()
            try:
                avail = int(avail_val)
                if avail < 0:
                    raise ValueError

                siblings     = get_siblings(book_code)
                total_copies = len(siblings)

                if avail > total_copies:
                    return render_edit(book=book,
                                        error=f"Available copies cannot exceed total copies ({total_copies}).")

                # Sort siblings by their copy number for deterministic assignment
                def copy_num(s):
                    try:
                        return int(s.code.rsplit("c.", 1)[-1])
                    except ValueError:
                        return 0
                siblings_sorted = sorted(siblings, key=copy_num)

                # First `avail` copies become available (1), rest unavailable (0)
                for i, sib in enumerate(siblings_sorted):
                    sib.available_copies = 1 if i < avail else 0

                db.session.commit()
                book = Book.query.filter_by(code=book_code).first()
                return render_edit(book=book,
                                    success=f"Availability updated: {avail}/{total_copies} copies marked available.")
            except ValueError:
                return render_edit(book=book,
                                    error="Available copies must be a non-negative integer.")
            except Exception as exc:
                db.session.rollback()
                return render_edit(book=book,
                                    error=f"Update failed: {exc}")

        # ══════════════════════════════════════════════════════════════════
        #  ADD COPY
        # ══════════════════════════════════════════════════════════════════
        elif action == "add_copy":
            parts = book_code.split("-")
            try:
                running_part = parts[-2]
                subject_part = parts[-3]
                main_part    = "-".join(parts[:-3])
                running_no   = int(running_part)
            except (IndexError, ValueError):
                return render_edit(book=book,
                                    error="Could not parse book code to add a copy.")

            siblings = get_siblings(book_code)
            max_copy = 0
            for s in siblings:
                try:
                    n = int(s.code.rsplit("c.", 1)[-1])
                    if n > max_copy:
                        max_copy = n
                except ValueError:
                    continue

            new_copy_no = max_copy + 1
            new_code    = f"{main_part}-{subject_part}-{running_part}-c.{new_copy_no}"

            if not is_code_unique(new_code):
                return render_edit(book=book,
                                    error=f"Copy code '{new_code}' already exists.")

            cover_file = request.files.get("new_copy_cover")
            if cover_file and cover_file.filename:
                ext         = os.path.splitext(secure_filename(cover_file.filename))[1].lower()
                unique_name = f"{new_code.replace('/', '_')}_{int(datetime.utcnow().timestamp())}{ext}"
                os.makedirs(COVER_DIR, exist_ok=True)
                cover_file.save(os.path.join(COVER_DIR, unique_name))
                cover_filename = unique_name
            else:
                cover_filename = book.cover_image

            try:
                new_total = len(siblings) + 1
                db.session.add(Book(
                    code             = new_code,
                    title            = book.title,
                    author           = book.author,
                    publisher        = book.publisher,
                    edition          = book.edition,
                    publication_year = book.publication_year,
                    subject_category = book.subject_category,
                    shelf_number     = book.shelf_number,
                    keywords         = book.keywords,
                    price            = book.price,
                    copies           = new_total,
                    available_copies = 1,
                    remarks          = book.remarks,
                    cover_image      = cover_filename,
                ))
                for s in siblings:
                    s.copies = new_total
                db.session.commit()
                book = Book.query.filter_by(code=book_code).first()
                return render_edit(book=book,
                                    success=f"New copy added as '{new_code}'.")
            except Exception as exc:
                db.session.rollback()
                return render_edit(book=book,
                                    error=f"Failed to add copy: {exc}")

        # ══════════════════════════════════════════════════════════════════
        #  EDIT (propagates to all siblings)
        # ══════════════════════════════════════════════════════════════════
        elif action == "edit":
            title = request.form.get("title", "").strip()
            if not title:
                return render_edit(book=book,
                                    error="Title cannot be empty.")

            def keep(field, old):
                v = request.form.get(field, "").strip()
                return v if v else old

            pub_val  = request.form.get("publication_year", "").strip()
            pub_year = int(float(pub_val)) if pub_val else book.publication_year

            price_val = request.form.get("price", "").strip()
            price_int = int(float(price_val)) if price_val else book.price

            subj_name = request.form.get("subject_name", "").strip().upper()
            if subj_name:
                subj = Subject.query.filter_by(name=subj_name).first()
                if not subj:
                    return render_edit(book=book,
                                        error=f"Subject '{subj_name}' not found.")
                subj_id = subj.id
            else:
                subj_id = book.subject_category

            new_shelf = request.form.get("shelf_number", "").strip() or book.shelf_number

            cover_file = request.files.get("cover_image")
            if cover_file and cover_file.filename:
                ext         = os.path.splitext(secure_filename(cover_file.filename))[1].lower()
                unique_name = f"{book_code.replace('/', '_')}_{int(datetime.utcnow().timestamp())}{ext}"
                os.makedirs(COVER_DIR, exist_ok=True)
                if book.cover_image and book.cover_image != "DEFAULTCOVER.png":
                    old_path = os.path.join(COVER_DIR, book.cover_image)
                    if os.path.exists(old_path):
                        os.remove(old_path)
                cover_file.save(os.path.join(COVER_DIR, unique_name))
                new_cover = unique_name
            else:
                new_cover = book.cover_image

            shared_fields = dict(
                title            = title,
                author           = keep("author",    book.author),
                publisher        = keep("publisher", book.publisher),
                edition          = keep("edition",   book.edition),
                publication_year = pub_year,
                subject_category = subj_id,
                shelf_number     = new_shelf,
                keywords         = keep("keywords",  book.keywords),
                price            = price_int,
                remarks          = keep("remarks",   book.remarks),
            )

            try:
                siblings = get_siblings(book_code)
                for sibling in siblings:
                    for field, value in shared_fields.items():
                        setattr(sibling, field, value)
                    if sibling.code == book_code:
                        sibling.cover_image = new_cover

                db.session.commit()
                book = Book.query.filter_by(code=book_code).first()
                sibling_codes = [s.code for s in siblings if s.code != book_code]
                extra = (f" Also updated {len(sibling_codes)} sibling cop{'y' if len(sibling_codes)==1 else 'ies'}: "
                         + ", ".join(sibling_codes)) if sibling_codes else ""
                return render_edit(book=book,
                                    success=f"Book '{book_code}' updated.{extra}")
            except Exception as exc:
                db.session.rollback()
                return render_edit(book=book,
                                    error=f"Update failed: {exc}")

    return render_edit(book=book, search_error=search_error)
# ── EXPORT BOOKS ──────────────────────────────

@admin_bp.route("/books/export", methods=["GET", "POST"])
@admin_required
def export_books():
    mains    = Main.query.order_by(Main.name).all()
    subjects = Subject.query.order_by(Subject.name).all()

    if request.method == "POST":
        fmt         = request.form.get("format")
        filter_type = request.form.get("filter_type")
        filter_id   = request.form.get("filter_id", "").strip()

        q = Book.query
        label = "All"
        if filter_type == "main" and filter_id:
            m = Main.query.get(int(filter_id))
            if m:
                q = q.filter(Book.code.ilike(f"{m.name}-%"))
                label = f"Main-{m.name}"
        elif filter_type == "subject" and filter_id:
            s = Subject.query.get(int(filter_id))
            if s:
                q = q.filter(Book.subject_category == s.id)
                label = f"Subj-{s.name}"

        all_books = q.order_by(Book.code).all()
        if not all_books:
            return render_template("admin/books/export.html",
                                   mains=mains, subjects=subjects,
                                   error="No books found for the selected filter.")

        HEADERS = ["Code","Title","Author","Publisher","Edition","Publication Year",
                   "Subject","Shelf No","Keywords","Price","Total Copies",
                   "Available Copies","Remarks","Added On"]

        def row(b):
            return [b.code, b.title, b.author or "", b.publisher or "",
                    b.edition or "", b.publication_year or "",
                    b.subject.name if b.subject else "",
                    b.shelf_number or "", b.keywords or "", b.price or "",
                    b.copies, b.available_copies, b.remarks or "",
                    b.created_at.strftime("%Y-%m-%d") if b.created_at else ""]

        fn = f"books_{label}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        if fmt == "csv":
            import csv
            out = io.StringIO()
            w   = csv.writer(out)
            w.writerow(HEADERS)
            for b in all_books: w.writerow(row(b))
            out.seek(0)
            return send_file(io.BytesIO(out.getvalue().encode()),
                mimetype="text/csv", as_attachment=True,
                download_name=f"{fn}.csv")

        elif fmt == "excel":
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Books"
            ws.append(HEADERS)
            for cell in ws[1]:
                cell.font  = Font(bold=True, color="FFFFFF")
                cell.fill  = PatternFill("solid", fgColor="1F4E79")
                cell.alignment = Alignment(horizontal="center")
            light = PatternFill("solid", fgColor="DCE6F1")
            for i, b in enumerate(all_books, start=2):
                ws.append(row(b))
                if i % 2 == 0:
                    for cell in ws[i]: cell.fill = light
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = (
                    min(max(len(str(c.value or "")) for c in col) + 4, 50))
            out = io.BytesIO(); wb.save(out); out.seek(0)
            return send_file(out,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True, download_name=f"{fn}.xlsx")

        elif fmt == "pdf":
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.platypus import (SimpleDocTemplate, Table,
                                            TableStyle, Paragraph, Spacer)
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors
            from reportlab.lib.units import cm
            out = io.BytesIO()
            doc = SimpleDocTemplate(out, pagesize=landscape(A4),
                                    leftMargin=1*cm, rightMargin=1*cm,
                                    topMargin=1.5*cm, bottomMargin=1.5*cm)
            styles = getSampleStyleSheet()
            s7     = styles["Normal"].clone("s7")
            s7.fontSize = 7; s7.leading = 9
            def wrap(v): return Paragraph(str(v), s7)
            elements = [
                Paragraph(f"Book Export — {label}", styles["Title"]),
                Paragraph(
                    f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}  |  "
                    f"Total: {len(all_books)}",
                    styles["Normal"]
                ),
                Spacer(1, 0.4*cm),
            ]
            td = [[wrap(h) for h in HEADERS]] + [[wrap(v) for v in row(b)] for b in all_books]
            cw = [w*cm for w in [3.2,4.5,3.0,3.0,1.8,1.5,2.2,1.8,3.5,1.2,1.2,1.5,3.0,1.8]]
            t  = Table(td, colWidths=cw, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F4E79")),
                ("TEXTCOLOR",(0,0),(-1,0),colors.white),
                ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
                ("FONTSIZE",(0,0),(-1,-1),7),
                ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#DCE6F1")]),
                ("GRID",(0,0),(-1,-1),0.4,colors.grey),
                ("VALIGN",(0,0),(-1,-1),"TOP"),
                ("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),
                ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
            ]))
            elements.append(t)
            doc.build(elements)
            out.seek(0)
            return send_file(out, mimetype="application/pdf",
                as_attachment=True, download_name=f"{fn}.pdf")

    return render_template("admin/books/export.html", mains=mains, subjects=subjects)


# ── CODE MAKER ────────────────────────────────

@admin_bp.route("/books/codes", methods=["GET", "POST"])
@admin_required
def book_codes():
    mains    = Main.query.order_by(Main.name).all()
    subjects = Subject.query.order_by(Subject.name).all()
    shelves  = Shelf.query.order_by(Shelf.name).all()

    if request.method == "POST":
        action = request.form.get("action")

        if action == "add_main":
            name = request.form.get("main_name", "").strip().upper()
            desc = request.form.get("main_desc", "").strip() or None
            if not name:
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    main_error="Name required.")
            if Main.query.filter_by(name=name).first():
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    main_error=f"'{name}' already exists.")
            try:
                db.session.add(Main(name=name, description=desc))
                db.session.commit()
                return redirect(url_for("admin.book_codes"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    main_error=str(exc))

        elif action == "delete_main":
            main_id = request.form.get("main_id", type=int)
            obj = Main.query.get(main_id)
            if not obj:
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    main_error="Main category not found.")
            if Book.query.filter(Book.code.ilike(f"{obj.name}-%")).first():
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    main_error=f"Cannot delete '{obj.name}' — books exist with this code.")
            try:
                db.session.delete(obj)
                db.session.commit()
                return redirect(url_for("admin.book_codes"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    main_error=str(exc))

        elif action == "add_subject":
            name = request.form.get("subject_name", "").strip().upper()
            desc = request.form.get("subject_desc", "").strip() or None
            if not name:
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    subject_error="Name required.")
            if Subject.query.filter_by(name=name).first():
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    subject_error=f"'{name}' already exists.")
            try:
                db.session.add(Subject(name=name, description=desc))
                db.session.commit()
                return redirect(url_for("admin.book_codes"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    subject_error=str(exc))

        elif action == "delete_subject":
            subj_id = request.form.get("subject_id", type=int)
            obj = Subject.query.get(subj_id)
            if not obj:
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    subject_error="Subject not found.")
            if Book.query.filter_by(subject_category=obj.id).first():
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    subject_error=f"Cannot delete '{obj.name}' — books are assigned to it.")
            try:
                db.session.delete(obj)
                db.session.commit()
                return redirect(url_for("admin.book_codes"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    subject_error=str(exc))

        elif action == "add_shelf":
            name = request.form.get("shelf_name", "").strip()
            desc = request.form.get("shelf_desc", "").strip() or None
            if not name:
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error="Shelf name required.")
            if Shelf.query.filter_by(name=name).first():
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error=f"'{name}' already exists.")
            try:
                db.session.add(Shelf(name=name, description=desc))
                db.session.commit()
                return redirect(url_for("admin.book_codes"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error=str(exc))

        elif action == "delete_shelf":
            shelf_id = request.form.get("shelf_id", type=int)
            obj = Shelf.query.get(shelf_id)
            if not obj:
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error="Shelf not found.")
            if Book.query.filter_by(shelf_number=obj.name).first():
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error=f"Cannot delete '{obj.name}' — books are assigned to it.")
            try:
                db.session.delete(obj)
                db.session.commit()
                return redirect(url_for("admin.book_codes"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error=str(exc))
        elif action == "edit_shelf":
            shelf_id = request.form.get("shelf_id", type=int)
            obj = Shelf.query.get(shelf_id)
            if not obj:
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error="Shelf not found.")
            new_desc = request.form.get("shelf_desc", "").strip() or None
            try:
                obj.description = new_desc
                db.session.commit()
                return redirect(url_for("admin.book_codes"))
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/books/codes.html",
                    mains=mains, subjects=subjects, shelves=shelves,
                    shelf_error=str(exc))
    return render_template("admin/books/codes.html",
                           mains=mains, subjects=subjects, shelves=shelves)


# ══════════════════════════════════════════════
#  EMPLOYEES SECTION
# ══════════════════════════════════════════════

@admin_bp.route("/employees")
@admin_required
def employees():
    tab       = request.args.get("tab",       "employees")
    search    = request.args.get("search",    "").strip().lower()
    search_by = request.args.get("search_by", "name")

    eq = Employee.query
    if search and tab == "employees":
        if search_by == "name":
            eq = eq.filter(Employee.name.ilike(f"%{search}%"))
        else:
            eq = eq.filter(Employee.employee_code.ilike(f"%{search}%"))
    all_employees = eq.order_by(Employee.name).all()
    employee_data = [{"employee": e, "indicator": get_employee_indicator(e)}
                     for e in all_employees]

    gq = GuestUser.query
    if search and tab == "guests":
        if search_by == "name":
            gq = gq.filter(GuestUser.name.ilike(f"%{search}%"))
        else:
            gq = gq.join(Employee, GuestUser.employee_code == Employee.employee_code)\
                   .filter(Employee.employee_code.ilike(f"%{search}%"))
    all_guests = gq.order_by(GuestUser.created_at.desc()).all()
    # Guests are listed with basic info only — no borrow indicator
    guest_data = [{"guest": g} for g in all_guests]

    return render_template(
        "admin/employees/index.html",
        tab=tab, employee_data=employee_data, guest_data=guest_data,
        search=search, search_by=search_by,
    )


@admin_bp.route("/employees/<string:employee_code>/detail")
@admin_required
def employee_detail(employee_code):
    employee = Employee.query.get_or_404(employee_code)
    now      = datetime.utcnow()
    all_records    = (Record.query.filter_by(employee_id=employee_code)
                      .order_by(Record.issued_at.desc()).all())
    total_taken    = len([r for r in all_records if r.status in ("APPROVED", "RETURNED")])
    total_returned = len([r for r in all_records if r.status == "RETURNED"])
    currently_have = [r for r in all_records if r.status == "APPROVED"]
    overdue        = [r for r in currently_have if r.expected_at and r.expected_at < now]
    roles          = Role.query.filter_by(employee_code=employee_code).all()
    return render_template("admin/employees/detail.html",
        employee=employee, all_records=all_records,
        total_taken=total_taken, total_returned=total_returned,
        currently_have=currently_have, overdue=overdue,
        indicator=get_employee_indicator(employee), roles=roles)


@admin_bp.route("/guests/<int:guest_id>/detail")
@admin_required
def guest_detail(guest_id):
    """Shows guest visitor information only — no borrow history."""
    guest = GuestUser.query.get_or_404(guest_id)
    sponsor = Employee.query.get(guest.employee_code)
    return render_template("admin/employees/guest_detail.html",
        guest=guest, sponsor=sponsor)


@admin_bp.route("/guests/manage", methods=["GET", "POST"])
@admin_required
def manage_guest():
    guest        = None
    search_error = None
    search_id    = request.args.get("id", "").strip()

    if search_id:
        try:
            guest = GuestUser.query.get(int(search_id))
            if not guest:
                search_error = f"No guest found with ID '{search_id}'."
        except (ValueError, TypeError):
            search_error = "Guest ID must be a number."

    if request.method == "POST":
        action = request.form.get("action")

        # ── ADD ───────────────────────────────────────────────────────────
        if action == "add":
            name           = request.form.get("name",           "").strip()
            contact_number = request.form.get("contact_number", "").strip() or None
            sponsor_code   = request.form.get("employee_code",  "").strip().upper()

            missing = []
            if not name:         missing.append("Name")
            if not sponsor_code: missing.append("Sponsor Employee Code")
            if missing:
                return render_template("admin/employees/manage_guest.html",
                                       guest=guest,
                                       error=f"Required: {', '.join(missing)}")

            sponsor = Employee.query.filter_by(
                employee_code=sponsor_code, status="ACTIVE"
            ).first()
            if not sponsor:
                return render_template("admin/employees/manage_guest.html",
                                       guest=guest,
                                       error=f"Sponsor '{sponsor_code}' not found or inactive.")

            try:
                new_guest = GuestUser(
                    name           = name,
                    employee_code  = sponsor_code,
                    contact_number = contact_number,
                )
                db.session.add(new_guest)
                db.session.commit()
                return render_template("admin/employees/manage_guest.html",
                                       guest=new_guest,
                                       success=f"Guest '{name}' added with ID #{new_guest.id}.")
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/employees/manage_guest.html",
                                       guest=guest,
                                       error=f"Could not add guest: {exc}")

        # ── EDIT ──────────────────────────────────────────────────────────
        elif action == "edit":
            guest_id = request.form.get("guest_id", type=int)
            guest    = GuestUser.query.get(guest_id)
            if not guest:
                return render_template("admin/employees/manage_guest.html",
                                       error=f"Guest #{guest_id} not found.")

            def keep(field, old):
                v = request.form.get(field, "").strip()
                return v if v else old

            new_name    = keep("name",           guest.name)
            new_contact = keep("contact_number", guest.contact_number)
            new_sponsor = request.form.get("employee_code", "").strip().upper()

            if new_sponsor and new_sponsor != guest.employee_code:
                sponsor = Employee.query.filter_by(
                    employee_code=new_sponsor, status="ACTIVE"
                ).first()
                if not sponsor:
                    return render_template("admin/employees/manage_guest.html",
                                           guest=guest,
                                           error=f"Sponsor '{new_sponsor}' not found or inactive.")
                guest.employee_code = new_sponsor

            guest.name           = new_name
            guest.contact_number = new_contact if new_contact else None

            try:
                db.session.commit()
                guest = GuestUser.query.get(guest_id)
                return render_template("admin/employees/manage_guest.html",
                                       guest=guest,
                                       success=f"Guest #{guest_id} updated.")
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/employees/manage_guest.html",
                                       guest=guest,
                                       error=str(exc))

        # ── DELETE ────────────────────────────────────────────────────────
        elif action == "delete":
            guest_id = request.form.get("guest_id", type=int)
            guest    = GuestUser.query.get(guest_id)
            if not guest:
                return render_template("admin/employees/manage_guest.html",
                                       error=f"Guest #{guest_id} not found.")
            name = guest.name
            try:
                db.session.delete(guest)
                db.session.commit()
                return render_template("admin/employees/manage_guest.html",
                                       success=f"Guest '{name}' (#{guest_id}) deleted.")
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/employees/manage_guest.html",
                                       guest=guest,
                                       error=f"Delete failed: {exc}")

    return render_template("admin/employees/manage_guest.html",
                           guest=guest, search_error=search_error)


@admin_bp.route("/employees/sponsor-lookup")
@admin_required
def sponsor_lookup():
    code = request.args.get("code", "").strip().upper()
    if not code:
        return jsonify({"found": False})

    employee = Employee.query.filter_by(
        employee_code=code, status="ACTIVE"
    ).first()

    if not employee:
        return jsonify({"found": False})

    return jsonify({
        "found":       True,
        "code":        employee.employee_code,
        "name":        employee.name,
        "department":  employee.department or "",
        "designation": employee.designation or "",
    })


@admin_bp.route("/employees/manage", methods=["GET", "POST"])
@admin_required
def manage_employee():
    employee     = None
    search_error = None
    search_code  = request.args.get("code", "").strip().lower()
    if search_code:
        employee = Employee.query.filter_by(employee_code=search_code).first()
        if not employee:
            search_error = f"No employee found with code '{search_code}'."

    if request.method == "POST":
        action = request.form.get("action")

        if action == "add":
            emp_code = request.form.get("employee_code", "").strip()
            name     = request.form.get("name",          "").strip()
            password = request.form.get("password",      "").strip()
            missing  = []
            if not emp_code: missing.append("Employee Code")
            if not name:     missing.append("Name")
            if not password: missing.append("Password")
            if missing:
                return render_template("admin/employees/manage.html",
                                       error=f"Required: {', '.join(missing)}")
            if Employee.query.get(emp_code):
                return render_template("admin/employees/manage.html",
                                       error=f"Code '{emp_code}' already exists.")

            gmail = request.form.get("gmail_id", "").strip() or None
            if gmail and Employee.query.filter_by(gmail_id=gmail).first():
                return render_template("admin/employees/manage.html",
                                       error=f"Gmail '{gmail}' is already registered to another employee.")

            role_name = request.form.get("role", "employee").strip().lower()
            try:
                db.session.add(Employee(
                    employee_code  = emp_code, name=name,
                    department     = request.form.get("department",     "").strip() or None,
                    designation    = request.form.get("designation",    "").strip() or None,
                    gmail_id       = gmail,
                    contact_number = request.form.get("contact_number", "").strip() or None,
                    status         = "ACTIVE",
                ))
                db.session.add(User(employee_code=emp_code, password_hash=p.hash(password)))
                db.session.add(Role(employee_code=emp_code, name=role_name))
                db.session.commit()
                return render_template("admin/employees/manage.html",
                                       success=f"Employee '{emp_code}' added.")
            except Exception as exc:
                db.session.rollback()
                msg = str(exc)
                if "gmail_id" in msg:
                    friendly = f"Gmail '{gmail}' is already registered to another employee."
                elif "employee_code" in msg or "employees_pkey" in msg:
                    friendly = f"Employee code '{emp_code}' already exists."
                elif "contact_number" in msg:
                    friendly = "That contact number is already registered."
                else:
                    friendly = "Could not add employee — please check your details and try again."
                return render_template("admin/employees/manage.html", error=friendly)

        elif action == "edit":
            emp_code = request.form.get("employee_code", "").strip().upper()
            employee = Employee.query.get(emp_code)
            if not employee:
                return render_template("admin/employees/manage.html",
                                       error=f"Employee '{emp_code}' not found.")
            def keep(field, old):
                v = request.form.get(field, "").strip()
                return v if v else old
            employee.name           = keep("name",           employee.name)
            employee.department     = keep("department",     employee.department)
            employee.designation    = keep("designation",    employee.designation)
            employee.gmail_id       = keep("gmail_id",       employee.gmail_id)
            employee.contact_number = keep("contact_number", employee.contact_number)
            new_status = request.form.get("status", "").strip().upper()
            if new_status in ("ACTIVE", "INACTIVE"):
                employee.status = new_status
            new_role = request.form.get("role", "").strip().lower()
            if new_role in ("employee", "admin"):
                Role.query.filter_by(employee_code=emp_code).delete()
                db.session.add(Role(employee_code=emp_code, name=new_role))
            try:
                db.session.commit()
                employee = Employee.query.get(emp_code)
                return render_template("admin/employees/manage.html",
                                       employee=employee, success=f"'{emp_code}' updated.")
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/employees/manage.html",
                                       employee=employee, error=str(exc))

        elif action == "delete":
            emp_code = request.form.get("employee_code", "").strip().upper()
            employee = Employee.query.get(emp_code)
            if not employee:
                return render_template("admin/employees/manage.html",
                                       error=f"Employee '{emp_code}' not found.")
            try:
                db.session.delete(employee)
                db.session.commit()
                return render_template("admin/employees/manage.html",
                                       success=f"Employee '{emp_code}' deleted.")
            except Exception as exc:
                db.session.rollback()
                return render_template("admin/employees/manage.html",
                                       employee=employee, error=str(exc))

    return render_template("admin/employees/manage.html",
                           employee=employee, search_error=search_error)


@admin_bp.route("/employees/records")
@admin_required
def employee_records():
    now           = datetime.utcnow()
    filter        = request.args.get("filter", "all")
    all_employees = Employee.query.order_by(Employee.name).all()
    overdue_list, active_list, free_list = [], [], []
    for e in all_employees:
        active_recs = [r for r in e.records if r.status == "APPROVED"]
        if not active_recs:
            free_list.append(e)
        elif any(r.expected_at and r.expected_at < now for r in active_recs):
            overdue_list.append(e)
        else:
            active_list.append(e)
    display = {"overdue": overdue_list, "active": active_list,
               "free": free_list}.get(filter, all_employees)
    if request.args.get("export") == "1":
        wb = openpyxl.Workbook(); ws = wb.active
        ws.title = filter.capitalize()
        ws.append(["Code","Name","Department","Designation","Contact","Status"])
        for e in display:
            ws.append([e.employee_code, e.name, e.department or "",
                       e.designation or "", e.contact_number or "", e.status])
        out = io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"employees_{filter}_{datetime.now().strftime('%Y%m%d')}.xlsx")
    return render_template("admin/employees/records.html",
        overdue_count=len(overdue_list), active_count=len(active_list),
        free_count=len(free_list), display=display, filter=filter,
        now=datetime.utcnow())


@admin_bp.route("/employees/approval/books")
@admin_required
def book_approval():
    pending         = Record.query.filter_by(status="PENDING").order_by(Record.issued_at).all()
    pending_returns = Record.query.filter_by(status="RETURN_REQUESTED").order_by(Record.returned_at).all()
    return render_template(
        "admin/employees/book_approval.html",
        pending=pending,
        pending_returns=pending_returns,
        now             = datetime.utcnow(),
    )


@admin_bp.route("/employees/approval/books/<int:record_id>/approve", methods=["POST"])
@admin_required
def approve_book(record_id):
    record = Record.query.get_or_404(record_id)
    if record.status == "PENDING":
        record.status      = "APPROVED"
        record.approval_at = datetime.utcnow()
        record.issued_at   = datetime.utcnow()

        # Use admin-edited date if provided, else keep employee's requested date,
        # else fall back to 14 days
        expected_date_str = request.form.get("expected_date", "").strip()
        if expected_date_str:
            try:
                record.expected_at = datetime.strptime(expected_date_str, "%Y-%m-%d")
            except ValueError:
                record.expected_at = datetime.utcnow() + timedelta(days=14)
        elif not record.expected_at:
            record.expected_at = datetime.utcnow() + timedelta(days=14)
        # else: keep the employee's originally requested date unchanged

        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
    return redirect(url_for("admin.book_approval"))


@admin_bp.route("/employees/approval/returns/<int:record_id>/approve", methods=["POST"])
@admin_required
def approve_return(record_id):
    record = Record.query.get_or_404(record_id)
    if record.status == "RETURN_REQUESTED":
        book = Book.query.get(record.book_id)
        if book:
            book.available_copies += 1
        record.status             = "RETURNED"
        record.return_approval_at = datetime.utcnow()
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
    return redirect(url_for("admin.book_approval"))


@admin_bp.route("/employees/approval/returns/<int:record_id>/reject", methods=["POST"])
@admin_required
def reject_return(record_id):
    record = Record.query.get_or_404(record_id)
    if record.status == "RETURN_REQUESTED":
        record.status      = "APPROVED"
        record.returned_at = None
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
    return redirect(url_for("admin.book_approval"))


@admin_bp.route("/employees/approval/books/<int:record_id>/reject", methods=["POST"])
@admin_required
def reject_book(record_id):
    record = Record.query.get_or_404(record_id)
    if record.status == "PENDING":
        book = Book.query.get(record.book_id)
        if book: book.available_copies += 1
        record.status = "REJECTED"
        try: db.session.commit()
        except Exception: db.session.rollback()
    return redirect(url_for("admin.book_approval"))


@admin_bp.route("/employees/approval/employees")
@admin_required
def employee_approval():
    pending = Employee.query.filter_by(status="INACTIVE").order_by(Employee.created_at).all()
    return render_template("admin/employees/employee_approval.html", pending=pending)


@admin_bp.route("/employees/approval/employees/<string:employee_code>/activate", methods=["POST"])
@admin_required
def activate_employee(employee_code):
    employee        = Employee.query.get_or_404(employee_code)
    employee.status = "ACTIVE"
    try: db.session.commit()
    except Exception: db.session.rollback()
    return redirect(url_for("admin.employee_approval"))


@admin_bp.route("/employees/approval/employees/<string:employee_code>/reject", methods=["POST"])
@admin_required
def reject_employee_registration(employee_code):
    employee = Employee.query.get_or_404(employee_code)
    try: db.session.delete(employee); db.session.commit()
    except Exception: db.session.rollback()
    return redirect(url_for("admin.employee_approval"))


@admin_bp.route("/employees/export", methods=["GET", "POST"])
@admin_required
def export_employees():
    if request.method == "POST":
        export_type = request.form.get("export_type", "all")
        now         = datetime.utcnow()
        if export_type == "guests":
            rows     = GuestUser.query.order_by(GuestUser.created_at).all()
            headers  = ["Guest Name","Sponsor Code","Sponsor Name","Contact","Registered On"]
            data     = [[g.name, g.employee_code, g.sponsor.name if g.sponsor else "",
                         g.contact_number or "",
                         g.created_at.strftime("%Y-%m-%d") if g.created_at else ""] for g in rows]
            filename = f"guests_{datetime.now().strftime('%Y%m%d')}.xlsx"
        else:
            q = Employee.query
            if export_type == "active":     q = q.filter_by(status="ACTIVE")
            elif export_type == "inactive": q = q.filter_by(status="INACTIVE")
            all_emp = q.order_by(Employee.name).all()
            if export_type == "overdue":
                all_emp = [e for e in all_emp if any(
                    r.status == "APPROVED" and r.expected_at and r.expected_at < now
                    for r in e.records)]
            headers  = ["Code","Name","Department","Designation","Gmail","Contact","Status","Joined"]
            data     = [[e.employee_code, e.name, e.department or "", e.designation or "",
                         e.gmail_id or "", e.contact_number or "", e.status,
                         e.created_at.strftime("%Y-%m-%d") if e.created_at else ""] for e in all_emp]
            filename = f"employees_{export_type}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center")
        for row in data: ws.append(row)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = (
                max(len(str(c.value or "")) for c in col) + 4)
        out = io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True, download_name=filename)
    return render_template("admin/employees/export.html")
@admin_bp.route("/employees/direct_issue", methods=["GET", "POST"])
@admin_required
def direct_issue():
    if request.method == 'POST':
        employee_code = g.current_user["employee_code"]
        name          = request.form.get("name",          "").strip()
        book_id       = request.form.get("book_id",       "").strip()
        expected_date = request.form.get("expected_date", "").strip()
        remarks       = request.form.get("remarks",       "").strip() or None   # ← new

        expected_at = datetime.strptime(expected_date, "%Y-%m-%d") if expected_date else None

        if not name or not book_id:
            return render_template("admin/employees/direct_issue.html",
                                   msg="Name and Book Code are mandatory.")

        if expected_at and expected_at.date() < datetime.utcnow().date():
            return render_template("admin/employees/direct_issue.html",
                                   msg="Expected return date cannot be in the past.")

        book = Book.query.filter(Book.code.ilike(book_id)).first()
        if not book:
            return render_template("admin/employees/direct_issue.html",
                                   msg=f"No book found with code '{book_id}'.")

        if book.available_copies < 1:
            return render_template("admin/employees/direct_issue.html",
                                   msg=f"'{book.code}' has no available copies.")

        try:
            from app.models import AdminRecord
            book.available_copies -= 1
            db.session.add(AdminRecord(
                employee_id = employee_code,
                book_id     = book.code,
                name        = name,
                issued_at   = datetime.utcnow(),
                expected_at = expected_at,
                status      = 'ISSUED',
                remarks     = remarks,             # ← new
            ))
            db.session.commit()
            return render_template("admin/employees/direct_issue.html",
                                   msg=f"Book '{book.code}' issued successfully.")
        except Exception as exc:
            db.session.rollback()
            return render_template("admin/employees/direct_issue.html",
                                   msg=f"Issue failed: {exc}")

    return render_template("admin/employees/direct_issue.html", msg='')

@admin_bp.route("/employees/adminrecord/<int:id>", methods=["GET", "POST"])
@admin_required
def adminrecord(id):
    details = AdminRecord.query.filter_by(id=id).first_or_404()

    if request.method == "POST":
        new_status = request.form.get("status", "").strip().upper()
        if new_status not in ("ISSUED", "RETURNED"):
            return render_template("admin/employees/adminrecords.html",
                                   msg="Invalid status value.", details=details)

        if new_status == "RETURNED" and details.status != "RETURNED":
            # ── increment available copies on the book ──
            book = Book.query.filter_by(code=details.book_id).first()
            if book:
                book.available_copies += 1
            else:
                return render_template("admin/employees/adminrecords.html",
                                       msg=f"Book '{details.book_id}' not found in DB.",
                                       details=details)
            details.returned_at = datetime.utcnow()

        # Guard against re-issuing an already-returned record
        elif new_status == "ISSUED" and details.status == "RETURNED":
            return render_template("admin/employees/adminrecords.html",
                                   msg="Cannot re-issue an already returned book from here.",
                                   details=details)

        details.status = new_status
        try:
            db.session.commit()
            return render_template("admin/employees/adminrecords.html",
                                   msg=f"Status updated to {new_status}.",
                                   details=details)
        except Exception as exc:
            db.session.rollback()
            return render_template("admin/employees/adminrecords.html",
                                   msg=f"Update failed: {exc}", details=details)

    return render_template("admin/employees/adminrecords.html", msg="", details=details)
@admin_bp.route("/employees/adminrecords")
@admin_required
def admin_records_list():
    records = AdminRecord.query.order_by(AdminRecord.issued_at.desc()).all()
    return render_template("admin/employees/adminrecord_list.html", records=records)
@admin_bp.route("/employees/adminrecords/delete-all", methods=["POST"])
@admin_required
def delete_all_admin_records():
    try:
        AdminRecord.query.delete(synchronize_session=False)
        db.session.commit()
    except Exception:
        db.session.rollback()
    return redirect(url_for("admin.admin_records_list"))